"""Read-only profiling and privacy-preserving linkage for clinical XLSX files.

Aggregate profiling never returns patient identifiers.  The explicit linkage
API uses identifiers only transiently and returns HMAC-pseudonymised keys; it
has no output-path or logging surface and never edits a source workbook.
"""

from __future__ import annotations

import hashlib
import hmac
import json
import re
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .privacy import is_placeholder_secret


class WorkbookDependencyError(RuntimeError):
    """Raised when XLSX support is requested without ``openpyxl`` installed."""


class WorkbookProfileError(ValueError):
    """Raised when a workbook cannot be safely profiled."""


@dataclass(frozen=True, slots=True)
class QuarantineSummary:
    """Aggregate-only quarantine metadata.

    Row numbers and identifier values are deliberately excluded because either
    can become a patient-level lookup key when combined with the source file.
    """

    missing_identifier_rows: int = 0
    duplicate_identifier_rows: int = 0
    duplicate_identifier_groups: int = 0

    @property
    def total_rows(self) -> int:
        return self.missing_identifier_rows + self.duplicate_identifier_rows


@dataclass(frozen=True, slots=True)
class SheetProfile:
    sheet_name: str
    row_count: int
    column_count: int
    headers: tuple[str, ...]
    identifier_header: str | None
    distinct_patient_count: int
    quarantine: QuarantineSummary


@dataclass(frozen=True, slots=True)
class WorkbookProfile:
    """Non-identifying metadata for one workbook."""

    source_document_id: str
    sheet_profiles: tuple[SheetProfile, ...]
    row_count: int
    distinct_patient_count: int
    quarantine: QuarantineSummary


@dataclass(frozen=True, slots=True)
class PatientSetCheck:
    """Aggregate patient-set relationship between two workbooks."""

    left_source_document_id: str
    right_source_document_id: str
    left_count: int
    right_count: int
    intersection_count: int
    union_count: int
    left_only_count: int
    right_only_count: int


@dataclass(frozen=True, slots=True)
class DataProfile:
    workbooks: tuple[WorkbookProfile, ...]
    total_rows: int
    summed_distinct_patient_count: int
    union_patient_count: int
    patients_in_all_workbooks: int
    patient_set_checks: tuple[PatientSetCheck, ...] = field(default_factory=tuple)


@dataclass(frozen=True, slots=True)
class QuarantinedRow:
    """Non-identifying pointer to a source row rejected from linkage.

    The row digest is a domain-separated HMAC, rather than a plain digest, so
    low-entropy identifiers cannot be recovered with an offline dictionary.
    By contract this object carries no row number, identifier, or source text.
    """

    source_document_id: str
    sheet_name: str
    reason: str
    row_hash: str


@dataclass(frozen=True, slots=True)
class SourceCell:
    """Location of one structured field in a read-only source workbook."""

    source_document_id: str
    sheet_name: str
    cell_ref: str


@dataclass(frozen=True, slots=True)
class PatientFieldObservation:
    """One non-identifier field retained in a linked patient view."""

    field_name: str
    value: Any
    source: SourceCell


@dataclass(frozen=True, slots=True)
class PatientFieldConflict:
    """References observations that disagree for one canonical field."""

    field_name: str
    observation_indexes: tuple[int, ...]
    distinct_value_count: int


@dataclass(frozen=True, slots=True)
class LinkedPatientView:
    """HMAC-pseudonymised patient-level view across source tables."""

    patient_key: str
    observations: tuple[PatientFieldObservation, ...]
    conflicts: tuple[PatientFieldConflict, ...] = field(default_factory=tuple)


@dataclass(frozen=True, slots=True)
class PatientViewLinkage:
    """In-memory linkage result; this module deliberately offers no writer."""

    patient_views: tuple[LinkedPatientView, ...]
    quarantine: tuple[QuarantinedRow, ...]
    source_document_count: int
    sheet_count: int


DEFAULT_IDENTIFIER_ALIASES: tuple[str, ...] = (
    "patient_id",
    "patientid",
    "患者id",
    "患者编号",
    "病人编号",
    "病案号",
    "住院号",
    "就诊号",
    "病例号",
)

DIRECT_IDENTIFIER_HEADERS: tuple[str, ...] = (
    *DEFAULT_IDENTIFIER_ALIASES,
    "患者姓名",
    "病人姓名",
    "姓名",
    "身份证号",
    "证件号码",
    "手机号",
    "联系电话",
    "电话",
    "电子邮箱",
    "email",
    "家庭住址",
    "地址",
)


def _normalise_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", "_", text)
    return text.strip("_-／/")


def _normalise_identifier(value: Any) -> str | None:
    """Normalise an identifier in memory without converting it to a number."""

    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text:
        return None
    # Excel sometimes renders integer identifiers as ``123.0``.
    if re.fullmatch(r"[0-9]+\.0", text):
        text = text[:-2]
    return text


def flatten_headers(
    rows: Sequence[Sequence[Any]], *, levels: int = 4
) -> tuple[str, ...]:
    """Flatten a multi-row XLSX header into stable, unique column names.

    Empty cells are horizontally forward-filled at each level to model merged
    group headings.  Repeated adjacent components are collapsed.  The function
    accepts fewer than four rows for tests and small synthetic templates, while
    production callers use the four-row default.
    """

    if levels < 1:
        raise ValueError("levels must be positive")
    selected = list(rows[:levels])
    if not selected:
        return ()
    width = max((len(row) for row in selected), default=0)
    if width == 0:
        return ()

    filled: list[list[str]] = []
    for raw_row in selected:
        normalised = [_normalise_header(v) for v in raw_row]
        normalised.extend([""] * (width - len(normalised)))
        carry = ""
        row: list[str] = []
        for value in normalised:
            if value:
                carry = value
            row.append(value or carry)
        filled.append(row)

    names: list[str] = []
    seen: Counter[str] = Counter()
    for column_index in range(width):
        components: list[str] = []
        for row in filled:
            value = row[column_index]
            if value and (not components or components[-1] != value):
                components.append(value)
        base = "__".join(components) or f"unnamed_{column_index + 1}"
        seen[base] += 1
        names.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return tuple(names)


def _source_document_id(path: Path) -> str:
    """Create a stable opaque source id without returning the path or filename."""

    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"xlsx_{digest.hexdigest()}"


def _identifier_column(
    headers: Sequence[str], aliases: Sequence[str]
) -> int | None:
    alias_set = {
        re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", alias.casefold())
        for alias in aliases
    }
    for index, header in enumerate(headers):
        components = header.split("__")
        for component in reversed(components):
            key = re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", component.casefold())
            if key in alias_set:
                return index
    return None


def _aggregate_quarantine(parts: Iterable[QuarantineSummary]) -> QuarantineSummary:
    parts = tuple(parts)
    return QuarantineSummary(
        missing_identifier_rows=sum(p.missing_identifier_rows for p in parts),
        duplicate_identifier_rows=sum(p.duplicate_identifier_rows for p in parts),
        duplicate_identifier_groups=sum(p.duplicate_identifier_groups for p in parts),
    )


def _load_openpyxl() -> Any:
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on installation extra
        raise WorkbookDependencyError(
            "XLSX profiling requires the optional 'openpyxl' dependency"
        ) from exc
    return openpyxl


@dataclass(frozen=True, slots=True)
class _LinkageRow:
    """Transient source row used only inside :func:`link_patient_views`."""

    row_number: int
    values: tuple[Any, ...]
    identifier: str
    row_hash: str


def _hmac_secret_bytes(secret: bytes | str) -> bytes:
    encoded = secret.encode("utf-8") if isinstance(secret, str) else secret
    if not isinstance(encoded, bytes) or len(encoded) < 16:
        raise WorkbookProfileError("hmac_secret must contain at least 16 bytes")
    if is_placeholder_secret(secret):
        raise WorkbookProfileError("hmac_secret cannot be a placeholder")
    return encoded


def _keyed_digest(secret: bytes, purpose: bytes, payload: bytes) -> str:
    return hmac.new(secret, purpose + b"\0" + payload, hashlib.sha256).hexdigest()


def _patient_key(secret: bytes, identifier: str) -> str:
    digest = _keyed_digest(secret, b"patient-key-v1", identifier.encode("utf-8"))
    return f"pt_{digest}"


def _digest_cell(value: Any) -> tuple[str, str]:
    if value is None:
        return ("null", "")
    if isinstance(value, (date, datetime)):
        return (type(value).__name__, value.isoformat())
    if isinstance(value, bytes):
        return ("bytes", value.hex())
    return (type(value).__name__, str(value))


def _row_hash(
    secret: bytes,
    *,
    source_document_id: str,
    sheet_name: str,
    row_number: int,
    values: Sequence[Any],
) -> str:
    # Source coordinates are included inside the MAC but not returned.  This
    # prevents equal low-entropy rows in different locations sharing a digest.
    canonical = json.dumps(
        {
            "source": source_document_id,
            "sheet": sheet_name,
            "row": row_number,
            "values": [_digest_cell(value) for value in values],
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return f"row_{_keyed_digest(secret, b'quarantine-row-v1', canonical)}"


def _header_key(value: str) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", value.casefold())


def _specified_identifier_column(
    headers: Sequence[str], identifier_header: str
) -> int | None:
    requested = _header_key(identifier_header)
    for index, header in enumerate(headers):
        if _header_key(header) == requested:
            return index
        if any(_header_key(component) == requested for component in header.split("__")):
            return index
    return None


def _canonical_field_name(
    header: str, field_aliases: Mapping[str, str]
) -> str:
    components = [component for component in header.split("__") if component]
    if components and components[-1].isdigit() and len(components) > 1:
        components.pop()
    leaf = components[-1] if components else header
    alias_lookup = {_header_key(alias): target for alias, target in field_aliases.items()}
    return alias_lookup.get(_header_key(leaf), leaf)


def _value_fingerprint(value: Any) -> tuple[str, str]:
    """Make clinically equivalent scalar representations compare consistently."""

    if isinstance(value, bool):
        return ("bool", str(value).lower())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        return ("number", str(int(numeric)) if numeric.is_integer() else repr(numeric))
    if isinstance(value, (date, datetime)):
        return ("date", value.isoformat())
    if isinstance(value, str):
        return ("text", re.sub(r"\s+", " ", value.strip()).casefold())
    return _digest_cell(value)


def _build_linked_patient_view(
    patient_key: str, observations: Sequence[PatientFieldObservation]
) -> LinkedPatientView:
    ordered = tuple(
        sorted(
            observations,
            key=lambda observation: (
                observation.field_name,
                observation.source.source_document_id,
                observation.source.sheet_name,
                observation.source.cell_ref,
            ),
        )
    )
    by_field: dict[str, list[int]] = {}
    for index, observation in enumerate(ordered):
        by_field.setdefault(observation.field_name, []).append(index)

    conflicts: list[PatientFieldConflict] = []
    for field_name, indexes in sorted(by_field.items()):
        fingerprints = {_value_fingerprint(ordered[index].value) for index in indexes}
        if len(fingerprints) > 1:
            conflicts.append(
                PatientFieldConflict(
                    field_name=field_name,
                    observation_indexes=tuple(indexes),
                    distinct_value_count=len(fingerprints),
                )
            )
    return LinkedPatientView(
        patient_key=patient_key,
        observations=ordered,
        conflicts=tuple(conflicts),
    )


def link_patient_views(
    paths: Sequence[str | Path],
    *,
    hmac_secret: bytes | str,
    identifier_header: str | None = None,
    identifier_aliases: Sequence[str] = DEFAULT_IDENTIFIER_ALIASES,
    field_aliases: Mapping[str, str] | None = None,
    header_rows: int = 4,
) -> PatientViewLinkage:
    """Link structured patient rows across XLSX tables without exposing IDs.

    Each worksheet is treated as an independent source table.  Missing keys,
    absent key columns, and *all* rows belonging to an intra-sheet duplicate
    key are quarantined rather than selected by position.  Identifier columns
    are excluded from observations.  The raw identifier exists only in local
    variables during this call and is replaced by a domain-separated HMAC key
    before a patient view is constructed.

    ``identifier_header`` is an optional explicit header (matching either a
    flattened header or one of its components).  When omitted, aliases are
    searched.  ``field_aliases`` maps source field names to a shared canonical
    name so semantically equivalent columns can participate in conflict checks.
    """

    if not paths:
        raise WorkbookProfileError("at least one workbook is required")
    if header_rows < 1:
        raise WorkbookProfileError("header_rows must be positive")
    secret = _hmac_secret_bytes(hmac_secret)
    aliases = field_aliases or {}
    direct_identifier_keys = {
        _header_key(item) for item in DIRECT_IDENTIFIER_HEADERS
    }
    openpyxl = _load_openpyxl()

    observations_by_patient: dict[str, list[PatientFieldObservation]] = {}
    quarantine: list[QuarantinedRow] = []
    source_document_ids: set[str] = set()
    sheet_count = 0

    for path in paths:
        source = Path(path)
        if source.suffix.casefold() not in {".xlsx", ".xlsm"}:
            raise WorkbookProfileError("only XLSX/XLSM workbooks can be linked")
        if not source.is_file():
            raise WorkbookProfileError(
                "workbook does not exist or is not a regular file"
            )

        source_document_id = _source_document_id(source)
        source_document_ids.add(source_document_id)
        workbook = openpyxl.load_workbook(
            filename=source,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            for sheet in workbook.worksheets:
                sheet_count += 1
                iterator = sheet.iter_rows(values_only=True)
                raw_headers: list[tuple[Any, ...]] = []
                for _ in range(header_rows):
                    try:
                        raw_headers.append(tuple(next(iterator)))
                    except StopIteration:
                        raw_headers.append(())
                headers = flatten_headers(raw_headers, levels=header_rows)
                identifier_index = (
                    _specified_identifier_column(headers, identifier_header)
                    if identifier_header is not None
                    else _identifier_column(headers, identifier_aliases)
                )

                accepted_candidates: list[_LinkageRow] = []
                for row_number, raw_row in enumerate(
                    iterator, start=header_rows + 1
                ):
                    row = tuple(raw_row)
                    if not any(
                        value is not None and str(value).strip() for value in row
                    ):
                        continue
                    digest = _row_hash(
                        secret,
                        source_document_id=source_document_id,
                        sheet_name=sheet.title,
                        row_number=row_number,
                        values=row,
                    )
                    if identifier_index is None:
                        quarantine.append(
                            QuarantinedRow(
                                source_document_id=source_document_id,
                                sheet_name=sheet.title,
                                reason="identifier_column_not_found",
                                row_hash=digest,
                            )
                        )
                        continue
                    identifier = (
                        _normalise_identifier(row[identifier_index])
                        if identifier_index < len(row)
                        else None
                    )
                    if identifier is None:
                        quarantine.append(
                            QuarantinedRow(
                                source_document_id=source_document_id,
                                sheet_name=sheet.title,
                                reason="missing_identifier",
                                row_hash=digest,
                            )
                        )
                        continue
                    accepted_candidates.append(
                        _LinkageRow(
                            row_number=row_number,
                            values=row,
                            identifier=identifier,
                            row_hash=digest,
                        )
                    )

                frequencies = Counter(row.identifier for row in accepted_candidates)
                for candidate in accepted_candidates:
                    if frequencies[candidate.identifier] > 1:
                        quarantine.append(
                            QuarantinedRow(
                                source_document_id=source_document_id,
                                sheet_name=sheet.title,
                                reason="duplicate_identifier",
                                row_hash=candidate.row_hash,
                            )
                        )
                        continue

                    patient_key = _patient_key(secret, candidate.identifier)
                    patient_observations = observations_by_patient.setdefault(
                        patient_key, []
                    )
                    width = max(len(headers), len(candidate.values))
                    for column_index in range(width):
                        if column_index == identifier_index:
                            continue
                        value = (
                            candidate.values[column_index]
                            if column_index < len(candidate.values)
                            else None
                        )
                        if value is None or (
                            isinstance(value, str) and not value.strip()
                        ):
                            continue
                        header = (
                            headers[column_index]
                            if column_index < len(headers)
                            else f"unnamed_{column_index + 1}"
                        )
                        header_components = header.split("__")
                        if any(
                            _header_key(component) in direct_identifier_keys
                            for component in header_components
                        ):
                            continue
                        field_name = _canonical_field_name(header, aliases)
                        patient_observations.append(
                            PatientFieldObservation(
                                field_name=field_name,
                                value=value,
                                source=SourceCell(
                                    source_document_id=source_document_id,
                                    sheet_name=sheet.title,
                                    cell_ref=(
                                        f"{openpyxl.utils.get_column_letter(column_index + 1)}"
                                        f"{candidate.row_number}"
                                    ),
                                ),
                            )
                        )
        finally:
            workbook.close()

    patient_views = tuple(
        _build_linked_patient_view(patient_key, observations)
        for patient_key, observations in sorted(observations_by_patient.items())
    )
    return PatientViewLinkage(
        patient_views=patient_views,
        quarantine=tuple(
            sorted(
                quarantine,
                key=lambda row: (
                    row.source_document_id,
                    row.sheet_name,
                    row.reason,
                    row.row_hash,
                ),
            )
        ),
        source_document_count=len(source_document_ids),
        sheet_count=sheet_count,
    )


def _profile_workbook_internal(
    path: str | Path,
    *,
    identifier_aliases: Sequence[str],
    header_rows: int,
) -> tuple[WorkbookProfile, set[str]]:
    source = Path(path)
    if source.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise WorkbookProfileError("only XLSX/XLSM workbooks can be profiled")
    if not source.is_file():
        raise WorkbookProfileError("workbook does not exist or is not a regular file")
    if header_rows < 1:
        raise WorkbookProfileError("header_rows must be positive")

    openpyxl = _load_openpyxl()
    # read_only avoids materialising/editing cells; data_only prevents formula
    # execution and returns only the cached value stored in the workbook.
    workbook = openpyxl.load_workbook(
        filename=source,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    source_id = _source_document_id(source)
    sheet_profiles: list[SheetProfile] = []
    workbook_patients: set[str] = set()
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            raw_headers: list[tuple[Any, ...]] = []
            for _ in range(header_rows):
                try:
                    raw_headers.append(tuple(next(iterator)))
                except StopIteration:
                    raw_headers.append(())
            headers = flatten_headers(raw_headers, levels=header_rows)
            id_index = _identifier_column(headers, identifier_aliases)

            identifiers: list[str] = []
            missing = 0
            rows = 0
            for raw_row in iterator:
                row = tuple(raw_row)
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                rows += 1
                if id_index is None or id_index >= len(row):
                    missing += 1
                    continue
                identifier = _normalise_identifier(row[id_index])
                if identifier is None:
                    missing += 1
                else:
                    identifiers.append(identifier)

            frequencies = Counter(identifiers)
            duplicate_groups = sum(1 for count in frequencies.values() if count > 1)
            # All rows in a duplicate group are quarantined; silently choosing a
            # "first" record would hide conflicting source information.
            duplicate_rows = sum(count for count in frequencies.values() if count > 1)
            patients = {key for key, count in frequencies.items() if count == 1}
            workbook_patients.update(patients)
            quarantine = QuarantineSummary(
                missing_identifier_rows=missing,
                duplicate_identifier_rows=duplicate_rows,
                duplicate_identifier_groups=duplicate_groups,
            )
            sheet_profiles.append(
                SheetProfile(
                    sheet_name=sheet.title,
                    row_count=rows,
                    column_count=len(headers),
                    headers=headers,
                    identifier_header=headers[id_index] if id_index is not None else None,
                    distinct_patient_count=len(patients),
                    quarantine=quarantine,
                )
            )
    finally:
        workbook.close()

    profile = WorkbookProfile(
        source_document_id=source_id,
        sheet_profiles=tuple(sheet_profiles),
        row_count=sum(sheet.row_count for sheet in sheet_profiles),
        distinct_patient_count=len(workbook_patients),
        quarantine=_aggregate_quarantine(sheet.quarantine for sheet in sheet_profiles),
    )
    return profile, workbook_patients


def profile_workbook(
    path: str | Path,
    *,
    identifier_aliases: Sequence[str] = DEFAULT_IDENTIFIER_ALIASES,
    header_rows: int = 4,
) -> WorkbookProfile:
    """Profile one workbook without exposing identifier sets."""

    profile, _ = _profile_workbook_internal(
        path,
        identifier_aliases=identifier_aliases,
        header_rows=header_rows,
    )
    return profile


def profile_workbooks(
    paths: Sequence[str | Path],
    *,
    identifier_aliases: Sequence[str] = DEFAULT_IDENTIFIER_ALIASES,
    header_rows: int = 4,
) -> DataProfile:
    """Profile workbooks and return aggregate cross-workbook set checks.

    The function never copies or saves a workbook and has no output-path
    argument by design.  Identifier sets remain local to this call and are
    discarded after aggregate counts are computed.
    """

    if not paths:
        raise WorkbookProfileError("at least one workbook is required")
    profiles: list[WorkbookProfile] = []
    patient_sets: list[set[str]] = []
    for path in paths:
        profile, patient_set = _profile_workbook_internal(
            path,
            identifier_aliases=identifier_aliases,
            header_rows=header_rows,
        )
        profiles.append(profile)
        patient_sets.append(patient_set)

    checks: list[PatientSetCheck] = []
    for left_index, left in enumerate(patient_sets):
        for right_index in range(left_index + 1, len(patient_sets)):
            right = patient_sets[right_index]
            intersection = left & right
            union = left | right
            checks.append(
                PatientSetCheck(
                    left_source_document_id=profiles[left_index].source_document_id,
                    right_source_document_id=profiles[right_index].source_document_id,
                    left_count=len(left),
                    right_count=len(right),
                    intersection_count=len(intersection),
                    union_count=len(union),
                    left_only_count=len(left - right),
                    right_only_count=len(right - left),
                )
            )

    union_patients = set().union(*patient_sets)
    all_patients = set.intersection(*patient_sets) if patient_sets else set()
    return DataProfile(
        workbooks=tuple(profiles),
        total_rows=sum(profile.row_count for profile in profiles),
        summed_distinct_patient_count=sum(
            profile.distinct_patient_count for profile in profiles
        ),
        union_patient_count=len(union_patients),
        patients_in_all_workbooks=len(all_patients),
        patient_set_checks=tuple(checks),
    )


def profile_to_dict(profile: DataProfile) -> Mapping[str, Any]:
    """Convert a profile to JSON-compatible aggregate metadata.

    Kept explicit rather than using ``asdict`` so future patient-level internal
    fields cannot accidentally become serialised.
    """

    return {
        "total_rows": profile.total_rows,
        "summed_distinct_patient_count": profile.summed_distinct_patient_count,
        "union_patient_count": profile.union_patient_count,
        "patients_in_all_workbooks": profile.patients_in_all_workbooks,
        "workbooks": [
            {
                "source_document_id": workbook.source_document_id,
                "row_count": workbook.row_count,
                "distinct_patient_count": workbook.distinct_patient_count,
                "quarantine": {
                    "missing_identifier_rows": workbook.quarantine.missing_identifier_rows,
                    "duplicate_identifier_rows": workbook.quarantine.duplicate_identifier_rows,
                    "duplicate_identifier_groups": workbook.quarantine.duplicate_identifier_groups,
                },
                "sheets": [
                    {
                        "sheet_name": sheet.sheet_name,
                        "row_count": sheet.row_count,
                        "column_count": sheet.column_count,
                        "headers": list(sheet.headers),
                        "identifier_header": sheet.identifier_header,
                        "distinct_patient_count": sheet.distinct_patient_count,
                        "quarantine": {
                            "missing_identifier_rows": sheet.quarantine.missing_identifier_rows,
                            "duplicate_identifier_rows": sheet.quarantine.duplicate_identifier_rows,
                            "duplicate_identifier_groups": sheet.quarantine.duplicate_identifier_groups,
                        },
                    }
                    for sheet in workbook.sheet_profiles
                ],
            }
            for workbook in profile.workbooks
        ],
        "patient_set_checks": [
            {
                "left_source_document_id": check.left_source_document_id,
                "right_source_document_id": check.right_source_document_id,
                "left_count": check.left_count,
                "right_count": check.right_count,
                "intersection_count": check.intersection_count,
                "union_count": check.union_count,
                "left_only_count": check.left_only_count,
                "right_only_count": check.right_only_count,
            }
            for check in profile.patient_set_checks
        ],
    }
